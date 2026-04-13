"""Excel导入导出工具"""
import openpyxl
from io import BytesIO
from typing import List, Dict, Any


def export_assets_to_excel(assets: List[Dict[str, Any]]) -> bytes:
    """导出资产数据到Excel"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "资产列表"

    # 表头
    headers = ["资产编号", "资产名称", "分类", "状态", "位置", "购置日期", "购置金额", "供应商"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    # 数据行
    for row, asset in enumerate(assets, 2):
        ws.cell(row=row, column=1, value=asset.get("asset_number"))
        ws.cell(row=row, column=2, value=asset.get("name"))
        ws.cell(row=row, column=3, value=asset.get("category_name", ""))
        ws.cell(row=row, column=4, value=asset.get("status"))
        ws.cell(row=row, column=5, value=asset.get("location"))
        ws.cell(row=row, column=6, value=str(asset.get("purchase_date", "")))
        ws.cell(row=row, column=7, value=float(asset.get("purchase_price", 0) or 0))
        ws.cell(row=row, column=8, value=asset.get("supplier"))

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def parse_assets_from_excel(file_content: bytes) -> List[Dict[str, Any]]:
    """从Excel解析资产数据"""
    wb = openpyxl.load_workbook(BytesIO(file_content))
    ws = wb.active

    assets = []
    headers = [cell.value for cell in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(row):
            asset = dict(zip(headers, row))
            assets.append(asset)

    return assets
